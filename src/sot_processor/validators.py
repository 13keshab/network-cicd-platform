"""Validation functions for the network CI/CD Excel Source of Truth."""

from collections import Counter
from ipaddress import ip_address, ip_interface, ip_network
from pathlib import Path

from openpyxl import load_workbook


WORKBOOK = (
    Path(__file__).resolve().parents[2]
    / "source_of_truth"
    / "network_source_of_truth.xlsx"
)

REQUIRED_SHEETS = {
    "Devices",
    "Interfaces",
    "Links",
    "Desired_Config",
    "Validation_Rules",
}

REQUIRED_COLUMNS = {
    "Devices": {
        "device_name", "environment", "role", "platform",
        "os_family", "connection", "management_ip", "enabled",
    },
    "Interfaces": {
        "device_name", "interface", "description", "ip_address",
        "prefix_length", "admin_state", "purpose",
    },
    "Links": {
        "link_id", "endpoint_a_device", "endpoint_a_interface",
        "endpoint_b_device", "endpoint_b_interface", "subnet",
        "environment", "link_type",
    },
    "Desired_Config": {
        "device_name", "feature", "setting", "value", "enabled", "notes",
    },
}


def read_rows(sheet):
    """Return non-empty worksheet rows as dictionaries."""
    headers = [cell.value for cell in sheet[1]]

    return [
        dict(zip(headers, values))
        for values in sheet.iter_rows(min_row=2, values_only=True)
        if any(value is not None for value in values)
    ]


def validate_workbook(workbook_path=WORKBOOK):
    """Return validation errors; an empty list means validation passed."""
    errors = []
    workbook_path = Path(workbook_path)

    if not workbook_path.is_file():
        return [f"Workbook not found: {workbook_path}"]

    try:
        workbook = load_workbook(
            workbook_path,
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        return [f"Cannot open workbook: {exc}"]

    try:
        # 1. Required worksheets and columns.
        missing_sheets = REQUIRED_SHEETS - set(workbook.sheetnames)

        if missing_sheets:
            return [
                f"Missing required sheets: {', '.join(sorted(missing_sheets))}"
            ]

        for sheet_name, required in REQUIRED_COLUMNS.items():
            headers = {
                cell.value
                for cell in workbook[sheet_name][1]
                if cell.value is not None
            }
            missing_columns = required - headers

            if missing_columns:
                errors.append(
                    f"{sheet_name}: missing required columns: "
                    f"{', '.join(sorted(missing_columns))}"
                )

        if errors:
            return errors

        devices = read_rows(workbook["Devices"])
        interfaces = read_rows(workbook["Interfaces"])
        links = read_rows(workbook["Links"])
        desired = read_rows(workbook["Desired_Config"])

        # 2. Device names must be present and unique.
        device_names = []

        for row in devices:
            name = row.get("device_name")

            if not isinstance(name, str) or not name.strip():
                errors.append("Device name is missing or invalid")
                continue

            device_names.append(name.strip())

        for name, count in Counter(device_names).items():
            if count > 1:
                errors.append(f"Duplicate device name: {name}")

        known_devices = set(device_names)

        # 3. Validate device management IPs.
        for row in devices:
            name = row.get("device_name")
            address = row.get("management_ip")

            if address is None or not str(address).strip():
                errors.append(f"Missing management IP for device {name}")
                continue

            try:
                parsed_address = ip_address(str(address).strip())
                if parsed_address.version != 4:
                    raise ValueError("IPv4 required")
            except ValueError:
                errors.append(
                    f"Invalid IPv4 management IP for {name}: {address!r}"
                )

        # 4. Validate interface IPs, prefixes, and duplicate IPs.
        interface_lookup = {}
        ip_owners = {}

        for row in interfaces:
            device = row.get("device_name")
            interface = row.get("interface")
            address = row.get("ip_address")
            prefix = row.get("prefix_length")

            if device not in known_devices:
                errors.append(
                    f"Interface {device}/{interface} references an unknown device"
                )
                continue

            if not isinstance(interface, str) or not interface.strip():
                errors.append(f"Missing interface name for device {device}")
                continue

            if address is None or prefix is None:
                errors.append(
                    f"Incomplete interface data for {device}/{interface}"
                )
                continue

            try:
                prefix_number = int(prefix)

                if not 0 <= prefix_number <= 32:
                    raise ValueError("Prefix must be between 0 and 32")

                parsed = ip_interface(
                    f"{str(address).strip()}/{prefix_number}"
                )

                if parsed.version != 4:
                    raise ValueError("IPv4 required")

            except (ValueError, TypeError):
                errors.append(
                    f"Invalid IPv4 address/prefix: "
                    f"{device}/{interface} ({address}/{prefix})"
                )
                continue

            key = (device, interface.strip())

            if key in interface_lookup:
                errors.append(
                    f"Duplicate interface definition: {device}/{interface}"
                )
                continue

            interface_lookup[key] = parsed

            ip = str(parsed.ip)

            if ip in ip_owners:
                errors.append(
                    f"Duplicate interface IP {ip}: "
                    f"{ip_owners[ip]} and {device}/{interface}"
                )
            else:
                ip_owners[ip] = f"{device}/{interface}"

        # 5. Link IDs must be present and unique.
        link_ids = []

        for row in links:
            link_id = row.get("link_id")

            if not isinstance(link_id, str) or not link_id.strip():
                errors.append("Link ID is missing or invalid")
                continue

            link_ids.append(link_id.strip())

        for link_id, count in Counter(link_ids).items():
            if count > 1:
                errors.append(f"Duplicate link ID: {link_id}")

        # 6. Validate links, endpoints, and subnet membership.
        for row in links:
            link_id = row.get("link_id") or "<unnamed-link>"
            a_device = row.get("endpoint_a_device")
            a_interface = row.get("endpoint_a_interface")
            b_device = row.get("endpoint_b_device")
            b_interface = row.get("endpoint_b_interface")
            subnet = row.get("subnet")

            try:
                network = ip_network(str(subnet), strict=True)

                if network.version != 4:
                    raise ValueError("IPv4 subnet required")

            except (ValueError, TypeError):
                errors.append(f"{link_id}: invalid subnet {subnet!r}")
                continue

            endpoints = (
                (a_device, a_interface),
                (b_device, b_interface),
            )
            endpoint_ips = []

            for device, interface in endpoints:
                # Ubuntu TAP is an external endpoint, not a router.
                if device == "Ubuntu-TAP" and interface == "tap-cicd":
                    if network != ip_network("10.10.100.0/24"):
                        errors.append(
                            f"{link_id}: TAP subnet should match "
                            "10.10.100.0/24"
                        )
                    continue

                if device not in known_devices:
                    errors.append(
                        f"{link_id}: unknown endpoint device {device!r}"
                    )
                    continue

                parsed = interface_lookup.get((device, interface))

                if parsed is None:
                    errors.append(
                        f"{link_id}: missing or invalid interface "
                        f"{device}/{interface}"
                    )
                    continue

                if parsed.ip not in network:
                    errors.append(
                        f"{link_id}: {device}/{interface} address "
                        f"{parsed.ip} is outside {network}"
                    )

                endpoint_ips.append(parsed.ip)

            if len(endpoint_ips) == 2 and endpoint_ips[0] == endpoint_ips[1]:
                errors.append(
                    f"{link_id}: both endpoints use the same IP "
                    f"{endpoint_ips[0]}"
                )


        # 7. Validate routing protocol intent.
        # R1 is shared infrastructure and is never an automation target.
        allowed_protocols = {"ospf", "eigrp", "rip", "ripv2"}
        routing_targets = {"R2", "R3"}

        for row in desired:
            if str(row.get("feature") or "").strip().lower() != "routing":
                continue

            protocol = str(row.get("setting") or "").strip().lower()
            device = str(row.get("device_name") or "").strip()
            value = str(row.get("value") or "").strip()
            enabled = row.get("enabled")

            if protocol not in allowed_protocols:
                errors.append(f"Unknown routing protocol {protocol!r} on {device}")
                continue
            if device not in known_devices:
                errors.append(f"Routing configuration references unknown device {device!r}")
                continue

            disabled = enabled is False or (
                isinstance(enabled, (int, float))
                and not isinstance(enabled, bool) and enabled == 0
            )
            enabled_ok = enabled is True or (
                isinstance(enabled, (int, float))
                and not isinstance(enabled, bool) and enabled == 1
            )

            if disabled:
                continue
            if not enabled_ok:
                errors.append(
                    f"Routing protocol {protocol} on {device} must have "
                    f"an explicit Boolean enabled value; got {enabled!r}"
                )
                continue
            if device not in routing_targets:
                errors.append(
                    f"Routing deployment is prohibited on {device}; "
                    "only R2 and R3 are permitted"
                )
                continue

            settings = {}
            for item in value.split(";"):
                if "=" not in item:
                    errors.append(f"Invalid {protocol.upper()} setting {item!r} on {device}")
                    continue
                key, setting_value = item.split("=", 1)
                key = key.strip().lower()
                setting_value = setting_value.strip()
                if not key or not setting_value or key in settings:
                    errors.append(
                        f"Empty or duplicate {protocol.upper()} setting "
                        f"{key!r} on {device}"
                    )
                    continue
                settings[key] = setting_value

            if protocol == "ospf":
                required = {"process_id", "area", "network"}
            elif protocol == "eigrp":
                required = {"as_number", "network"}
            else:
                required = {"version", "network"}

            missing = required - set(settings)
            unknown = set(settings) - required
            if missing:
                errors.append(
                    f"Missing {protocol.upper()} settings on {device}: "
                    f"{', '.join(sorted(missing))}"
                )
            if unknown:
                errors.append(
                    f"Unknown {protocol.upper()} settings on {device}: "
                    f"{', '.join(sorted(unknown))}"
                )
            if missing or unknown:
                continue

            if protocol == "ospf":
                try:
                    if not 1 <= int(settings["process_id"]) <= 65535:
                        raise ValueError
                except ValueError:
                    errors.append(f"Invalid OSPF process_id on {device}")
                try:
                    area = settings["area"]
                    if "." in area:
                        if ip_address(area).version != 4:
                            raise ValueError
                    elif not 0 <= int(area) <= 4294967295:
                        raise ValueError
                except ValueError:
                    errors.append(f"Invalid OSPF area on {device}")

            elif protocol == "eigrp":
                try:
                    if not 1 <= int(settings["as_number"]) <= 65535:
                        raise ValueError
                except ValueError:
                    errors.append(f"Invalid EIGRP AS number on {device}")

            elif settings["version"] != "2":
                errors.append(f"RIP on {device} must use version=2")

            network_texts = [
                item.strip()
                for item in settings["network"].split(",")
                if item.strip()
            ]
            if not network_texts:
                errors.append(f"No {protocol.upper()} networks specified on {device}")
                continue

            transit_networks = set()
            device_interfaces = []

            for interface_row in interfaces:
                if str(interface_row.get("device_name") or "").strip() != device:
                    continue
                try:
                    iface = ip_interface(
                        f"{str(interface_row['ip_address']).strip()}/"
                        f"{int(interface_row['prefix_length'])}"
                    )
                except (ValueError, TypeError, KeyError):
                    continue

                purpose = str(interface_row.get("purpose") or "").strip().lower()
                device_interfaces.append((iface.network, purpose))
                if purpose in {"staging_transit", "production_transit"}:
                    transit_networks.add(iface.network)

            for network_text in network_texts:
                try:
                    if protocol in {"ospf", "eigrp"}:
                        network = ip_network(network_text, strict=True)
                        if network.version != 4:
                            raise ValueError("IPv4 network required")
                    else:
                        address = ip_address(network_text)
                        if address.version != 4:
                            raise ValueError("IPv4 required")
                        first = int(network_text.split(".")[0])
                        if not 1 <= first <= 223:
                            raise ValueError("Invalid classful IPv4 network")
                        prefix = 8 if first <= 126 else 16 if first <= 191 else 24
                        network = ip_network(f"{address}/{prefix}", strict=False)
                        if network.network_address != address:
                            raise ValueError("Use the classful major network address")
                except ValueError as exc:
                    errors.append(
                        f"Invalid {protocol.upper()} network on {device}: "
                        f"{network_text!r} ({exc})"
                    )
                    continue

                if protocol in {"ospf", "eigrp"}:
                    if not any(
                        network.subnet_of(transit)
                        or transit.subnet_of(network)
                        for transit in transit_networks
                    ):
                        errors.append(
                            f"{protocol.upper()} network {network} on {device} "
                            "does not match a defined transit subnet"
                        )
                else:
                    # RIP network statements can match multiple interfaces.
                    unsafe = [
                        str(iface_network)
                        for iface_network, purpose in device_interfaces
                        if (
                            iface_network.network_address in network
                            or network.network_address in iface_network
                        )
                        and purpose not in {
                            "staging_transit", "production_transit"
                        }
                    ]
                    if unsafe:
                        errors.append(
                            f"RIP network {network} on {device} also matches "
                            "non-transit interface network(s): "
                            f"{', '.join(unsafe)}"
                        )


    finally:
        workbook.close()

    return errors


def main():
    errors = validate_workbook()

    if errors:
        print("SOURCE OF TRUTH VALIDATION FAILED")
        for error in errors:
            print(f" - {error}")
        raise SystemExit(1)

    print("SOURCE OF TRUTH VALIDATION PASSED")
    print(" - Required worksheets and columns exist")
    print(" - Device names and management IPs are checked")
    print(" - Interface IPs and prefixes are checked")
    print(" - Link IDs, endpoints, and subnets are checked")
    print(" - Routing protocol intent and enabled settings are validated")


if __name__ == "__main__":
    main()
