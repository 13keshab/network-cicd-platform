from openpyxl import load_workbook

from src.sot_processor.validators import WORKBOOK
from scripts.export_routing_intent import export_routing_intent


def test_current_workbook_exports_nested_routing_intent():
    result = export_routing_intent(WORKBOOK)
    routing = result["routing_intent"]

    assert set(routing) == {"R2", "R3"}
    assert "R1" not in routing

    # R2: OSPF is enabled in the current Source of Truth.
    r2_ospf = routing["R2"]["ospf"]
    assert r2_ospf["enabled"] is True
    assert r2_ospf["process_id"] == 10
    assert r2_ospf["area"] == "0"
    assert r2_ospf["networks"] == ["10.10.12.0/30"]
    assert r2_ospf["network_statements"] == [
        "network 10.10.12.0 0.0.0.3 area 0"
    ]

    # EIGRP and RIP are enabled on R2 in the current Source of Truth.
    assert routing["R2"]["eigrp"]["enabled"] is True
    assert routing["R2"]["eigrp"]["as_number"] == 100
    assert routing["R2"]["rip"]["enabled"] is True
    assert routing["R2"]["rip"]["version"] == 2

    # R3 has EIGRP and RIP enabled, but OSPF remains disabled.
    assert routing["R3"]["ospf"]["enabled"] is False
    assert routing["R3"]["eigrp"]["enabled"] is True
    assert routing["R3"]["eigrp"]["as_number"] == 100
    assert routing["R3"]["rip"]["enabled"] is True
    assert routing["R3"]["rip"]["version"] == 2

    assert routing["R3"]["ospf"]["process_id"] is None
    assert routing["R3"]["ospf"]["area"] is None
    assert routing["R3"]["ospf"]["networks"] == []


def test_enabled_ospf_is_exported_for_r2_only(tmp_path):
    path = tmp_path / "enabled_ospf.xlsx"
    workbook = load_workbook(WORKBOOK)
    sheet = workbook["Desired_Config"]

    for row in range(2, sheet.max_row + 1):
        if (
            sheet.cell(row, 1).value == "R2"
            and str(sheet.cell(row, 2).value or "").strip().lower()
            == "routing"
            and str(sheet.cell(row, 3).value or "").strip().lower()
            == "ospf"
        ):
            sheet.cell(row, 4).value = (
                "process_id=10; area=0; network=10.10.12.0/30"
            )
            sheet.cell(row, 5).value = True
            break
    else:
        workbook.close()
        raise AssertionError("Could not find the R2 OSPF configuration row")

    workbook.save(path)
    workbook.close()

    result = export_routing_intent(path)
    routing = result["routing_intent"]

    assert routing["R2"]["ospf"]["enabled"] is True
    assert routing["R2"]["ospf"]["process_id"] == 10
    assert routing["R2"]["ospf"]["area"] == "0"
    assert routing["R2"]["ospf"]["networks"] == ["10.10.12.0/30"]

    assert routing["R3"]["ospf"]["enabled"] is False
    assert routing["R3"]["ospf"]["process_id"] is None
    assert routing["R3"]["ospf"]["area"] is None
    assert routing["R3"]["ospf"]["networks"] == []

    # The shared router is never included in deployment intent.
    assert "R1" not in routing



def test_enabled_eigrp_is_exported_with_network_statements(tmp_path):
    path = tmp_path / "enabled_eigrp.xlsx"
    workbook = load_workbook(WORKBOOK)
    sheet = workbook["Desired_Config"]

    for row in range(2, sheet.max_row + 1):
        if (
            sheet.cell(row, 1).value == "R2"
            and str(sheet.cell(row, 2).value or "").strip().lower()
            == "routing"
            and str(sheet.cell(row, 3).value or "").strip().lower()
            == "eigrp"
        ):
            sheet.cell(row, 4).value = (
                "as_number=100; network=10.10.12.0/30"
            )
            sheet.cell(row, 5).value = True
            break
    else:
        workbook.close()
        raise AssertionError("Could not find the R2 EIGRP row")

    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 1).value == "R3" and str(sheet.cell(row, 3).value or "").strip().lower() == "eigrp":
            sheet.cell(row, 5).value = False
    workbook.save(path)
    workbook.close()

    result = export_routing_intent(path)["routing_intent"]

    assert result["R2"]["eigrp"]["enabled"] is True
    assert result["R2"]["eigrp"]["as_number"] == 100
    assert result["R2"]["eigrp"]["networks"] == ["10.10.12.0/30"]
    assert result["R2"]["eigrp"]["network_statements"] == [
        "network 10.10.12.0 0.0.0.3"
    ]
    assert result["R3"]["eigrp"]["enabled"] is False
    assert "R1" not in result


def test_enabled_ripv2_is_exported_with_network_statements(tmp_path):
    path = tmp_path / "enabled_rip.xlsx"
    workbook = load_workbook(WORKBOOK)
    sheet = workbook["Desired_Config"]

    for row in range(2, sheet.max_row + 1):
        if (
            sheet.cell(row, 1).value == "R2"
            and str(sheet.cell(row, 2).value or "").strip().lower()
            == "routing"
            and str(sheet.cell(row, 3).value or "").strip().lower()
            == "rip"
        ):
            sheet.cell(row, 4).value = (
                "version=2; network=10.0.0.0"
            )
            sheet.cell(row, 5).value = True
            break
    else:
        workbook.close()
        raise AssertionError("Could not find the R2 RIP row")

    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 1).value == "R3" and str(sheet.cell(row, 3).value or "").strip().lower() == "rip":
            sheet.cell(row, 5).value = False
    workbook.save(path)
    workbook.close()

    result = export_routing_intent(path)["routing_intent"]

    assert result["R2"]["rip"]["enabled"] is True
    assert result["R2"]["rip"]["version"] == 2
    assert result["R2"]["rip"]["networks"] == ["10.0.0.0"]
    assert result["R2"]["rip"]["network_statements"] == [
        "network 10.0.0.0"
    ]
    assert result["R3"]["rip"]["enabled"] is False
    assert "R1" not in result
