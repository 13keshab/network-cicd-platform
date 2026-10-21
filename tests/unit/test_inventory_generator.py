import yaml

from openpyxl import load_workbook

from src.sot_processor.inventory_generator import (
    build_inventory,
    write_inventories,
)
from src.sot_processor.validators import WORKBOOK


def test_combined_inventory_has_expected_groups():
    inventory = build_inventory(WORKBOOK)
    children = inventory["all"]["children"]

    assert set(children) == {"shared", "staging", "production"}
    assert set(children["shared"]["hosts"]) == {"R1"}
    assert set(children["staging"]["hosts"]) == {"R2"}
    assert set(children["production"]["hosts"]) == {"R3"}


def test_environment_inventories_are_written(tmp_path):
    inventory = build_inventory(WORKBOOK)

    write_inventories(inventory, tmp_path)

    for filename in (
        "all.yml",
        "shared.yml",
        "staging.yml",
        "production.yml",
    ):
        assert (tmp_path / filename).is_file()


def test_staging_inventory_contains_only_r2(tmp_path):
    inventory = build_inventory(WORKBOOK)
    write_inventories(inventory, tmp_path)

    data = yaml.safe_load((tmp_path / "staging.yml").read_text())

    hosts = data["all"]["children"]["staging"]["hosts"]
    assert set(hosts) == {"R2"}


def test_production_inventory_contains_only_r3(tmp_path):
    inventory = build_inventory(WORKBOOK)
    write_inventories(inventory, tmp_path)

    data = yaml.safe_load((tmp_path / "production.yml").read_text())

    hosts = data["all"]["children"]["production"]["hosts"]
    assert set(hosts) == {"R3"}


def test_disabled_device_is_excluded(tmp_path):
    test_workbook = tmp_path / "disabled_device.xlsx"
    workbook = load_workbook(WORKBOOK)

    sheet = workbook["Devices"]

    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 1).value == "R2":
            sheet.cell(row, 8).value = False
            break

    workbook.save(test_workbook)
    workbook.close()

    inventory = build_inventory(test_workbook)
    children = inventory["all"]["children"]

    assert "R2" not in children["staging"]["hosts"]
    assert "R1" in children["shared"]["hosts"]
    assert "R3" in children["production"]["hosts"]


def test_invalid_workbook_is_rejected(tmp_path):
    test_workbook = tmp_path / "invalid_workbook.xlsx"
    workbook = load_workbook(WORKBOOK)
    workbook["Interfaces"]["D2"] = "999.999.999.999"
    workbook.save(test_workbook)
    workbook.close()

    try:
        build_inventory(test_workbook)
    except ValueError as exc:
        assert "Source of Truth validation failed" in str(exc)
    else:
        raise AssertionError("Invalid workbook should not generate inventory")
