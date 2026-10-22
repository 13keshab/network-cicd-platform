from pathlib import Path

from openpyxl import load_workbook

from src.sot_processor.validators import WORKBOOK, validate_workbook


def make_test_workbook(tmp_path: Path, filename: str) -> Path:
    """Create an independent copy of the real workbook for a test."""
    test_path = tmp_path / filename
    workbook = load_workbook(WORKBOOK)
    workbook.save(test_path)
    workbook.close()
    return test_path


def test_current_source_of_truth_passes():
    assert validate_workbook(WORKBOOK) == []


def test_missing_workbook_fails(tmp_path):
    errors = validate_workbook(tmp_path / "missing.xlsx")

    assert errors
    assert any("Workbook not found" in error for error in errors)


def test_duplicate_device_name_fails(tmp_path):
    path = make_test_workbook(tmp_path, "duplicate_devices.xlsx")
    workbook = load_workbook(path)
    workbook["Devices"].append([
        "R1", "staging", "duplicate", "Cisco 7200VXR", "ios",
        "network_cli", "10.10.100.3", True
    ])
    workbook.save(path)
    workbook.close()

    errors = validate_workbook(path)

    assert any("Duplicate device name" in error for error in errors)


def test_duplicate_interface_ip_fails(tmp_path):
    path = make_test_workbook(tmp_path, "duplicate_ip.xlsx")
    workbook = load_workbook(path)
    workbook["Interfaces"].append([
        "R2", "GigabitEthernet2/0", "TEST-DUPLICATE",
        "10.10.12.2", 30, "up", "test"
    ])
    workbook.save(path)
    workbook.close()

    errors = validate_workbook(path)

    assert any("Duplicate interface IP" in error for error in errors)


def test_invalid_interface_ip_fails(tmp_path):
    path = make_test_workbook(tmp_path, "invalid_ip.xlsx")
    workbook = load_workbook(path)
    workbook["Interfaces"].append([
        "R2", "GigabitEthernet2/0", "TEST-INVALID",
        "999.10.12.2", 30, "up", "test"
    ])
    workbook.save(path)
    workbook.close()

    errors = validate_workbook(path)

    assert errors, "An invalid IPv4 address should be rejected"


def test_invalid_prefix_length_fails(tmp_path):
    path = make_test_workbook(tmp_path, "invalid_prefix.xlsx")
    workbook = load_workbook(path)
    workbook["Interfaces"].append([
        "R2", "GigabitEthernet2/0", "TEST-INVALID",
        "10.10.12.3", 33, "up", "test"
    ])
    workbook.save(path)
    workbook.close()

    errors = validate_workbook(path)

    assert errors, "An IPv4 prefix length greater than 32 should be rejected"


def test_link_with_wrong_subnet_fails(tmp_path):
    path = make_test_workbook(tmp_path, "wrong_subnet.xlsx")
    workbook = load_workbook(path)
    workbook["Links"]["F2"] = "192.0.2.0/30"
    workbook.save(path)
    workbook.close()

    errors = validate_workbook(path)

    assert errors, "A link subnet that does not contain its endpoint IPs should fail"


def test_missing_required_worksheet_fails(tmp_path):
    path = make_test_workbook(tmp_path, "missing_sheet.xlsx")
    workbook = load_workbook(path)
    workbook.remove(workbook["Links"])
    workbook.save(path)
    workbook.close()

    errors = validate_workbook(path)

    assert errors, "The workbook should fail if the Links worksheet is missing"



def test_valid_enabled_ospf_on_r2_passes(tmp_path):
    path = make_test_workbook(tmp_path, "valid_ospf_r2.xlsx")
    workbook = load_workbook(path)
    sheet = workbook["Desired_Config"]

    for row in range(2, sheet.max_row + 1):
        device = sheet.cell(row, 1).value
        feature = str(sheet.cell(row, 2).value or "").strip().lower()
        setting = str(sheet.cell(row, 3).value or "").strip().lower()

        if device == "R2" and feature == "routing" and setting == "ospf":
            sheet.cell(row, 4).value = (
                "process_id=10; area=0; network=10.10.12.0/30"
            )
            sheet.cell(row, 5).value = True
            break

    workbook.save(path)
    workbook.close()

    assert validate_workbook(path) == []


def test_enabled_ospf_on_r1_is_rejected(tmp_path):
    path = make_test_workbook(tmp_path, "ospf_r1.xlsx")
    workbook = load_workbook(path)
    sheet = workbook["Desired_Config"]

    for row in range(2, sheet.max_row + 1):
        if (
            sheet.cell(row, 1).value == "R1"
            and str(sheet.cell(row, 2).value or "").lower() == "routing"
            and str(sheet.cell(row, 3).value or "").lower() == "ospf"
        ):
            sheet.cell(row, 5).value = True
            break

    workbook.save(path)
    workbook.close()

    errors = validate_workbook(path)
    assert any("prohibited on R1" in error for error in errors)


def test_enabled_eigrp_is_rejected(tmp_path):
    path = make_test_workbook(tmp_path, "enabled_eigrp.xlsx")
    workbook = load_workbook(path)
    sheet = workbook["Desired_Config"]

    for row in range(2, sheet.max_row + 1):
        if (
            sheet.cell(row, 1).value == "R2"
            and str(sheet.cell(row, 2).value or "").lower() == "routing"
            and str(sheet.cell(row, 3).value or "").lower() == "eigrp"
        ):
            sheet.cell(row, 5).value = True
            break

    workbook.save(path)
    workbook.close()

    errors = validate_workbook(path)
    assert any("not implemented" in error for error in errors)


def test_invalid_enabled_ospf_network_is_rejected(tmp_path):
    path = make_test_workbook(tmp_path, "invalid_ospf_network.xlsx")
    workbook = load_workbook(path)
    sheet = workbook["Desired_Config"]

    for row in range(2, sheet.max_row + 1):
        if (
            sheet.cell(row, 1).value == "R2"
            and str(sheet.cell(row, 2).value or "").lower() == "routing"
            and str(sheet.cell(row, 3).value or "").lower() == "ospf"
        ):
            sheet.cell(row, 4).value = (
                "process_id=10; area=0; network=192.0.2.0/24"
            )
            sheet.cell(row, 5).value = True
            break

    workbook.save(path)
    workbook.close()

    errors = validate_workbook(path)
    assert any("does not match a defined transit subnet" in error
               for error in errors)


def test_duplicate_link_id_fails(tmp_path):
    path = make_test_workbook(tmp_path, "duplicate_link.xlsx")
    workbook = load_workbook(path)
    sheet = workbook["Links"]
    sheet.append([sheet.cell(2, column).value for column in range(1, 9)])
    workbook.save(path)
    workbook.close()

    errors = validate_workbook(path)

    assert errors, "Duplicate link IDs should be rejected"


def test_unknown_device_interface_reference_fails(tmp_path):
    path = make_test_workbook(tmp_path, "unknown_endpoint.xlsx")
    workbook = load_workbook(path)
    workbook["Links"]["B2"] = "R99"
    workbook.save(path)
    workbook.close()

    errors = validate_workbook(path)

    assert errors, "A link referencing an unknown device should be rejected"
